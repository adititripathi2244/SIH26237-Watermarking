"""
DOCX watermark carried by character spacing (Member 3's method).

Bit format is identical to Member 3's DOCXWatermarker, so his extractor can
read files made here and the other way round:

    MAGIC "VSM1" | version (1) | timestamp (8, big endian) | id length (1)
    | id bytes | CRC32 of everything before it (4)

    bit 1 -> <w:spacing w:val="10"/> on a one-character run
    bit 0 -> <w:spacing w:val="0"/>

Differences from the original, needed for the backend:
  * The caller passes the finished watermark_id (SHA-256 of the ZKP proof).
    Nothing is hashed here and no recipient identity ever reaches this module.
  * Run formatting (bold, italic, font, size, colour) is copied onto every
    one-character run. The original deleted all runs and re-created them
    plain, which wiped formatting and any inline picture in that paragraph.
  * Only plain-text runs carry bits, and only as many characters as the
    payload needs. Everything after that stays exactly as it was.
  * <w:spacing> is inserted at its schema position inside <w:rPr>, so Word
    does not see an out-of-order property.
  * Not enough text raises DocxCapacityError instead of a half-marked file.
"""

from __future__ import annotations

import copy
import struct
import zlib
from io import BytesIO

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

MAGIC = b"VSM1"
VERSION = 1
NORMAL_SPACING = "0"
WATERMARK_SPACING = "10"
MAX_ID_BYTES = 64
HEADER_BYTES = 14  # magic 4 + version 1 + timestamp 8 + id length 1

# rPr children that must come AFTER <w:spacing> in the OOXML schema.
_AFTER_SPACING = {
    qn(f"w:{name}")
    for name in (
        "w", "kern", "position", "sz", "szCs", "highlight", "u", "effect",
        "bdr", "shd", "fitText", "vertAlign", "rtl", "cs", "em", "lang",
        "eastAsianLayout", "specVanish", "oMath",
    )
}
_PLAIN_RUN_CHILDREN = {qn("w:rPr"), qn("w:t")}


class DocxCapacityError(ValueError):
    """The document has too little plain text to carry the watermark."""


# ------------------------------------------------------------------ payload

def _payload_bits(watermark_id: str, timestamp: int) -> list[int]:
    id_bytes = watermark_id.encode("ascii")
    if not id_bytes or len(id_bytes) > MAX_ID_BYTES:
        raise ValueError("watermark_id must be 1-64 ASCII characters")
    if not 0 <= timestamp <= 0xFFFFFFFFFFFFFFFF:
        raise ValueError("timestamp out of range")

    body = (
        MAGIC
        + struct.pack(">B", VERSION)
        + struct.pack(">Q", timestamp)
        + struct.pack(">B", len(id_bytes))
        + id_bytes
    )
    body += struct.pack(">I", zlib.crc32(body))
    return [(byte >> pos) & 1 for byte in body for pos in range(7, -1, -1)]


def _bits_to_bytes(bits: list[int]) -> bytes:
    out = bytearray()
    for i in range(0, len(bits) - 7, 8):
        value = 0
        for bit in bits[i:i + 8]:
            value = (value << 1) | bit
        out.append(value)
    return bytes(out)


# ------------------------------------------------------------------ runs

def _is_plain_text_run(run_element) -> bool:
    return all(child.tag in _PLAIN_RUN_CHILDREN for child in run_element)


def _run_text(run_element) -> str:
    return "".join(t.text or "" for t in run_element.findall(qn("w:t")))


def _set_spacing(run_element, bit: int) -> None:
    rpr = run_element.get_or_add_rPr()
    for old in rpr.findall(qn("w:spacing")):
        rpr.remove(old)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:val"), WATERMARK_SPACING if bit else NORMAL_SPACING)
    for child in rpr:
        if child.tag in _AFTER_SPACING:
            child.addprevious(spacing)
            break
    else:
        rpr.append(spacing)


def _clone_with_text(run_element, text: str):
    clone = copy.deepcopy(run_element)
    for child in list(clone):
        if child.tag != qn("w:rPr"):
            clone.remove(child)
    t = OxmlElement("w:t")
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    t.text = text
    clone.append(t)
    return clone


def _plain_runs(document):
    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            if _is_plain_text_run(run._r) and _run_text(run._r):
                yield run._r


def capacity(docx_bytes: bytes) -> int:
    """Number of characters available to carry bits."""
    document = Document(BytesIO(docx_bytes))
    return sum(len(_run_text(r)) for r in _plain_runs(document))


# ------------------------------------------------------------------ embed

def embed_watermark_id(docx_bytes: bytes, watermark_id: str, timestamp: int) -> bytes:
    bits = _payload_bits(watermark_id, int(timestamp))
    document = Document(BytesIO(docx_bytes))

    # Snapshot first: we change the tree while walking it.
    runs = list(_plain_runs(document))
    available = sum(len(_run_text(r)) for r in runs)
    if available < len(bits):
        raise DocxCapacityError(
            f"document has {available} usable characters, watermark needs {len(bits)}"
        )

    index = 0
    for run_element in runs:
        if index >= len(bits):
            break
        text = _run_text(run_element)
        take = min(len(text), len(bits) - index)

        parent = run_element.getparent()
        for char in text[:take]:
            carrier = _clone_with_text(run_element, char)
            _set_spacing(carrier, bits[index])
            run_element.addprevious(carrier)
            index += 1

        if take < len(text):
            run_element.addprevious(_clone_with_text(run_element, text[take:]))
        parent.remove(run_element)

    output = BytesIO()
    document.save(output)
    return output.getvalue()


# ------------------------------------------------------------------ extract

def _spacing_bits(document, count: int) -> list[int]:
    bits: list[int] = []
    for run_element in _plain_runs(document):
        rpr = run_element.rPr
        spacing = rpr.find(qn("w:spacing")) if rpr is not None else None
        value = spacing.get(qn("w:val")) if spacing is not None else None
        bit = 1 if value == WATERMARK_SPACING else 0
        for _ in _run_text(run_element):
            bits.append(bit)
            if len(bits) >= count:
                return bits
    return bits


def extract_watermark_id(docx_bytes: bytes) -> tuple[str, int]:
    """Return (watermark_id, timestamp). Raises ValueError if absent/corrupt."""
    document = Document(BytesIO(docx_bytes))

    head = _spacing_bits(document, HEADER_BYTES * 8)
    if len(head) < HEADER_BYTES * 8:
        raise ValueError("not enough text to hold a watermark")
    header = _bits_to_bytes(head)
    if header[:4] != MAGIC:
        raise ValueError("watermark signature not found")

    id_length = header[13]
    if not 1 <= id_length <= MAX_ID_BYTES:
        raise ValueError("watermark length is invalid")

    total = HEADER_BYTES + id_length + 4
    all_bits = _spacing_bits(document, total * 8)
    if len(all_bits) < total * 8:
        raise ValueError("watermark is truncated")
    payload = _bits_to_bytes(all_bits)

    if struct.unpack(">I", payload[-4:])[0] != zlib.crc32(payload[:-4]):
        raise ValueError("watermark checksum failed")

    timestamp = struct.unpack(">Q", payload[5:13])[0]
    try:
        watermark_id = payload[HEADER_BYTES:HEADER_BYTES + id_length].decode("ascii")
    except UnicodeDecodeError:
        raise ValueError("watermark id is not valid text") from None
    return watermark_id, timestamp
