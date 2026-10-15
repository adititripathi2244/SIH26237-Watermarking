"""
DOCX watermark tests (character-spacing mark from Member 3).

Run:  python -m pytest tests/api/test_docx.py -v
"""

import base64
import io

import pytest
from docx import Document
from docx.shared import Pt

from src.provenance import docx_watermarker as dw


def real_docx_bytes(clauses: int = 14) -> bytes:
    """A formatted Word file. clauses=0 gives a very short document."""
    doc = Document()
    doc.add_heading("Confidential Contract", 1)
    para = doc.add_paragraph()
    bold = para.add_run("Bold intro. ")
    bold.bold = True
    bold.font.size = Pt(14)
    para.add_run("Plain tail.").italic = True
    for n in range(clauses):
        doc.add_paragraph(
            f"Clause {n}: the recipient shall not disclose this document to any third party. " * 2
        )
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _decrypt(client, owner, file_id, private, who="alice"):
    return client.post("/decrypt", headers=owner, json={
        "encrypted_file_id": file_id, "recipient_id": who, "private_key": private[who]})


# ---------------------------------------------------------------- module level

def test_round_trip_and_formatting_is_kept():
    source = real_docx_bytes()
    marked = dw.embed_watermark_id(source, "a" * 64, 1759165200)

    assert dw.extract_watermark_id(marked) == ("a" * 64, 1759165200)

    before, after = Document(io.BytesIO(source)), Document(io.BytesIO(marked))
    assert [p.text for p in before.paragraphs] == [p.text for p in after.paragraphs]
    assert any(run.bold for run in after.paragraphs[1].runs)
    assert any(run.italic for run in after.paragraphs[1].runs)


def test_unmarked_document_has_no_watermark():
    with pytest.raises(ValueError):
        dw.extract_watermark_id(real_docx_bytes())


def test_too_short_document_is_refused_not_half_marked():
    with pytest.raises(dw.DocxCapacityError):
        dw.embed_watermark_id(real_docx_bytes(0), "a" * 64, 1)


# ---------------------------------------------------------------- through the API

def test_docx_full_loop(client, new_user, make_recipients, encrypt_doc):
    owner = new_user()
    public, private = make_recipients(["alice"])
    file_id = encrypt_doc(owner, public, filename="contract.docx", content=real_docx_bytes())

    r = _decrypt(client, owner, file_id, private)
    assert r.status_code == 200, r.text
    wid = r.json()["watermark_id"]
    assert len(wid) == 64                       # SHA-256(ZKP proof), same rule as PDF/PNG
    leaked = base64.b64decode(r.json()["watermarked_file"])

    ex = client.post("/extract-watermark", headers=owner, files={"file": ("leak.docx", leaked)})
    assert ex.status_code == 200, ex.text
    assert ex.json()["watermark_id"] == wid and ex.json()["confidence"] == 1.0

    tr = client.post("/provenance/trace", headers=owner, files={"file": ("leak.docx", leaked)})
    assert tr.status_code == 200, tr.text
    assert tr.json()["known"] is True and tr.json()["status"] == "confirmed"


def test_docx_other_user_cannot_trace(client, new_user, make_recipients, encrypt_doc):
    owner, intruder = new_user(), new_user()
    public, private = make_recipients(["alice"])
    file_id = encrypt_doc(owner, public, filename="contract.docx", content=real_docx_bytes())
    leaked = base64.b64decode(_decrypt(client, owner, file_id, private).json()["watermarked_file"])

    r = client.post("/provenance/trace", headers=intruder, files={"file": ("leak.docx", leaked)})
    assert r.status_code in (200, 404)
    if r.status_code == 200:
        assert r.json()["known"] is False
        assert r.json()["encrypted_file_id"] is None and r.json()["document_hash"] is None


def test_short_docx_falls_back_to_the_metadata_mark(client, new_user, make_recipients, encrypt_doc):
    owner = new_user()
    public, private = make_recipients(["alice"])
    file_id = encrypt_doc(owner, public, filename="short.docx", content=real_docx_bytes(0))

    r = _decrypt(client, owner, file_id, private)
    assert r.status_code == 200, r.text
    wid = r.json()["watermark_id"]
    assert len(wid) == 43                       # ForensicWatermark ID (fragile fallback)
    leaked = base64.b64decode(r.json()["watermarked_file"])

    ex = client.post("/extract-watermark", headers=owner, files={"file": ("leak.docx", leaked)})
    assert ex.status_code == 200, ex.text
    assert ex.json()["watermark_id"] == wid


def test_plain_docx_is_not_attributed(client, new_user):
    r = client.post("/provenance/trace", headers=new_user(),
                    files={"file": ("plain.docx", real_docx_bytes())})
    assert r.status_code in (404, 422) or (r.status_code == 200 and r.json()["known"] is False)
