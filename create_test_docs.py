from docx import Document


for i in range(1, 6):

    document = Document()

    document.add_heading(
        f"SIH Test Document {i}",
        level=1
    )

    document.add_paragraph(
        f"This is test document number {i} "
        "for DOCX watermarking. "
        "This document contains enough text characters "
        "to provide sufficient capacity for the watermark. "
        "The watermark will store recipient identification, "
        "timestamp information, version information, "
        "and checksum data for verification. "
        "This test document is part of the Phase 1 "
        "DOCX watermarking accuracy evaluation."
    )

    document.add_paragraph(
        "The purpose of this document is to test reliable "
        "embedding and extraction through the unified "
        "MultiFormatWatermarker interface. "
        "The same process will be repeated across five "
        "different DOCX documents."
    )

    document.add_paragraph(
        "Smart India Hackathon 2026 watermarking test. "
        "Recipient attribution and watermark extraction "
        "are being evaluated using the DOCX prototype."
    )

    document.save(f"test_doc_{i}.docx")

    print(f"Created test_doc_{i}.docx")