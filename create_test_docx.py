from docx import Document

# Create a new Word document
document = Document()

# Add a heading
document.add_heading(
    "SIH Watermarking Test Document",
    level=1
)

# Add some paragraphs
document.add_paragraph(
    "This is a test document for our forensic watermarking "
    "system developed for Smart India Hackathon 2026."
)

document.add_paragraph(
    "The purpose of this document is to test DOCX watermark "
    "embedding and extraction using character spacing."
)

document.add_paragraph(
    "This document contains enough text so that we can "
    "embed our recipient identification information."
)

document.add_paragraph(
    "Later, this document will also be used for robustness "
    "testing including compression, conversion, screenshot "
    "and print-scan simulation."
)

# Save the document
document.save("sample.docx")

print("sample.docx created successfully!")